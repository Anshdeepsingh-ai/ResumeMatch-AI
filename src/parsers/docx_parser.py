"""
DOCX Document Parser using python-docx.
Course: CSE472 Deep Learning for NLP
Milestone: Day 2 Resume Parser Module
"""

import zipfile
from pathlib import Path
from typing import Dict, Tuple

try:
    import docx
    from docx.opc.exceptions import OpcError, PackageNotFoundError
except ImportError:
    docx = None
    OpcError = Exception
    PackageNotFoundError = Exception

from src.parsers.exceptions import CorruptFileError, EmptyDocumentError, ParserError


def extract_docx_text(file_path: Path) -> Tuple[str, Dict[str, int]]:
    """
    Extract structured textual content from a DOCX document using python-docx.
    Extracts content from both body paragraphs and tabular layouts.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Tuple containing:
            - text (str): Extracted text content preserving paragraph boundaries.
            - metadata (Dict[str, int]): Technical metadata including paragraph and table counts.

    Raises:
        EmptyDocumentError: If the document contains no extractable text or paragraphs.
        CorruptFileError: If the DOCX package is corrupted, unreadable, or invalid.
    """
    if docx is None:
        raise ParserError("python-docx is not installed in the current environment.")

    try:
        doc = docx.Document(str(file_path))
    except (zipfile.BadZipFile, PackageNotFoundError, OpcError, KeyError, Exception) as exc:
        raise CorruptFileError(f"DOCX file is corrupt or invalid ({file_path.name}): {exc}") from exc

    extracted_parts = []
    paragraph_count = 0
    table_count = len(doc.tables)

    # 1. Extract standard paragraphs
    for p in doc.paragraphs:
        clean_text = p.text.strip()
        if clean_text:
            extracted_parts.append(clean_text)
            paragraph_count += 1

    # 2. Extract tabular text (commonly used in multi-column resume layouts)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            # De-duplicate identical merged cells in the same row
            unique_cells = []
            for cell_content in row_cells:
                if not unique_cells or cell_content != unique_cells[-1]:
                    unique_cells.append(cell_content)
            if unique_cells:
                extracted_parts.append(" | ".join(unique_cells))

    full_text = "\n\n".join(extracted_parts).strip()

    if not full_text:
        raise EmptyDocumentError(f"DOCX document contains no extractable text: {file_path.name}")

    metadata = {
        "paragraph_count": paragraph_count,
        "table_count": table_count,
        "total_extracted_blocks": len(extracted_parts),
    }

    return full_text, metadata
