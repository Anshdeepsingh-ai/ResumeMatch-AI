"""
Unit and integration tests for the ResumeMatch AI document parsing module.
Course: CSE472 Deep Learning for NLP
Milestone: Day 2 Resume Parser Module
"""

from pathlib import Path

import docx
import pytest

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from src.parsers import (
    CorruptFileError,
    DocumentParser,
    EmptyDocumentError,
    EncodingError,
    FileNotFoundParserError,
    FileTooLargeError,
    ParsedDocument,
    UnsupportedFormatError,
    extract_docx_text,
    extract_pdf_text,
    extract_txt_text,
    parse_document,
)


# ==============================================================================
# Helper Functions to Generate Synthetic Test Documents
# ==============================================================================


def create_synthetic_txt(path: Path, content: str, encoding: str = "utf-8") -> Path:
    """Create a synthetic text document for testing."""
    if encoding == "utf-8":
        path.write_text(content, encoding="utf-8")
    else:
        path.write_bytes(content.encode(encoding))
    return path


def create_synthetic_docx(path: Path, paragraphs=None, tables=None) -> Path:
    """Create a synthetic DOCX document for testing."""
    doc = docx.Document()
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if tables:
        for table_data in tables:
            t = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            for r_idx, row in enumerate(table_data):
                for c_idx, val in enumerate(row):
                    t.cell(r_idx, c_idx).text = val
    doc.save(str(path))
    return path


def create_synthetic_pdf(path: Path, pages_content=None) -> Path:
    """Create a synthetic PDF document for testing."""
    doc = fitz.open()
    if pages_content:
        for text in pages_content:
            page = doc.new_page()
            if text:
                page.insert_text((50, 72), text)
    else:
        doc.new_page()
    doc.save(str(path))
    doc.close()
    return path


# ==============================================================================
# Test Cases: Valid Document Extraction
# ==============================================================================


class TestValidDocumentExtraction:
    """Tests verifying correct text extraction across supported formats."""

    def test_extract_valid_txt(self, tmp_path):
        """Test extraction from a valid UTF-8 text file."""
        sample_text = (
            "Profile Summary\n"
            "Synthetic Candidate with 5 years experience in Natural Language Processing.\n\n"
            "Technical Skills\n"
            "Python, PyTorch, Transformers, SpaCy, Docker."
        )
        txt_file = create_synthetic_txt(tmp_path / "resume.txt", sample_text)

        result = parse_document(txt_file)
        assert isinstance(result, ParsedDocument)
        assert result.file_type == "txt"
        assert result.file_name == "resume.txt"
        assert "Synthetic Candidate" in result.text
        assert "Technical Skills" in result.text
        assert result.word_count > 10
        assert result.character_count == len(result.text)
        assert result.page_count is None

    def test_extract_valid_docx(self, tmp_path):
        """Test extraction from a valid DOCX file with paragraphs and tables."""
        paragraphs = [
            "Synthetic Candidate",
            "Role: Machine Learning Engineer",
            "Work Experience: Developed neural retrieval pipelines.",
        ]
        tables = [
            [["Skill", "Proficiency"], ["Python", "Expert"], ["PyTorch", "Advanced"]]
        ]
        docx_file = create_synthetic_docx(
            tmp_path / "resume.docx", paragraphs=paragraphs, tables=tables
        )

        result = parse_document(docx_file)
        assert isinstance(result, ParsedDocument)
        assert result.file_type == "docx"
        assert result.file_name == "resume.docx"
        assert "Synthetic Candidate" in result.text
        assert "Machine Learning Engineer" in result.text
        assert "Skill | Proficiency" in result.text or "Python | Expert" in result.text
        assert result.metadata["paragraph_count"] == 3
        assert result.metadata["table_count"] == 1

    def test_extract_valid_pdf(self, tmp_path):
        """Test extraction from a multi-page PDF document."""
        pages = [
            "Synthetic Candidate - Page 1\nSenior NLP Researcher\nSpecializing in Transformer Encoders.",
            "Synthetic Candidate - Page 2\nEducation: MS in Computational Linguistics.",
        ]
        pdf_file = create_synthetic_pdf(tmp_path / "resume.pdf", pages_content=pages)

        result = parse_document(pdf_file)
        assert isinstance(result, ParsedDocument)
        assert result.file_type == "pdf"
        assert result.file_name == "resume.pdf"
        assert result.page_count == 2
        assert "Synthetic Candidate - Page 1" in result.text
        assert "Synthetic Candidate - Page 2" in result.text
        assert result.metadata["non_empty_pages"] == 2

    def test_case_insensitive_extension(self, tmp_path):
        """Test that uppercase extensions (.PDF, .DOCX, .TXT) are parsed successfully."""
        pdf_file = create_synthetic_pdf(tmp_path / "TEST.PDF", ["Synthetic text in uppercase PDF."])
        docx_file = create_synthetic_docx(tmp_path / "TEST.DOCX", ["Synthetic text in uppercase DOCX."])
        txt_file = create_synthetic_txt(tmp_path / "TEST.TXT", "Synthetic text in uppercase TXT.")

        assert parse_document(pdf_file).file_type == "pdf"
        assert parse_document(docx_file).file_type == "docx"
        assert parse_document(txt_file).file_type == "txt"


# ==============================================================================
# Test Cases: Empty and Whitespace-Only Documents
# ==============================================================================


class TestEmptyAndWhitespaceDocuments:
    """Tests verifying proper detection of empty files and whitespace-only documents."""

    def test_empty_zero_byte_txt(self, tmp_path):
        """Zero-byte TXT file should raise EmptyDocumentError."""
        empty_file = tmp_path / "zero.txt"
        empty_file.touch()
        with pytest.raises(EmptyDocumentError, match="empty \\(0 bytes\\)"):
            parse_document(empty_file)

    def test_empty_zero_byte_pdf(self, tmp_path):
        """Zero-byte PDF file should raise EmptyDocumentError."""
        empty_file = tmp_path / "zero.pdf"
        empty_file.touch()
        with pytest.raises(EmptyDocumentError, match="empty \\(0 bytes\\)"):
            parse_document(empty_file)

    def test_empty_zero_byte_docx(self, tmp_path):
        """Zero-byte DOCX file should raise EmptyDocumentError."""
        empty_file = tmp_path / "zero.docx"
        empty_file.touch()
        with pytest.raises(EmptyDocumentError, match="empty \\(0 bytes\\)"):
            parse_document(empty_file)

    def test_whitespace_only_txt(self, tmp_path):
        """TXT file with spaces, tabs, and newlines should raise EmptyDocumentError."""
        ws_file = create_synthetic_txt(tmp_path / "whitespace.txt", "   \n\t  \n   \n")
        with pytest.raises(EmptyDocumentError, match="contains no extractable text"):
            parse_document(ws_file)

    def test_whitespace_only_docx(self, tmp_path):
        """DOCX file with blank paragraphs should raise EmptyDocumentError."""
        ws_file = create_synthetic_docx(tmp_path / "whitespace.docx", paragraphs=["  ", "\t\t", ""])
        with pytest.raises(EmptyDocumentError, match="contains no extractable text"):
            parse_document(ws_file)

    def test_scanned_or_textless_pdf(self, tmp_path):
        """PDF with pages but no extractable text (scanned document simulation) should raise EmptyDocumentError with OCR notice."""
        # Create a PDF with blank pages (no text layer)
        pdf_file = create_synthetic_pdf(tmp_path / "scanned_dummy.pdf", pages_content=["", "   "])
        with pytest.raises(EmptyDocumentError, match="OCR"):
            parse_document(pdf_file)


# ==============================================================================
# Test Cases: Corrupt and Invalid Documents
# ==============================================================================


class TestCorruptAndInvalidDocuments:
    """Tests verifying handling of corrupted files and encoding issues."""

    def test_corrupt_pdf(self, tmp_path):
        """Corrupt PDF file should raise CorruptFileError."""
        corrupt_file = tmp_path / "corrupt.pdf"
        corrupt_file.write_bytes(b"%PDF-1.4 header without valid xref trailer \x00\xff\xfe")
        with pytest.raises(CorruptFileError, match="corrupt or invalid"):
            parse_document(corrupt_file)

    def test_corrupt_docx(self, tmp_path):
        """Corrupt DOCX file (invalid zip archive) should raise CorruptFileError."""
        corrupt_file = tmp_path / "corrupt.docx"
        corrupt_file.write_bytes(b"PK\x03\x04 fake corrupted docx archive content \xff\xfe")
        with pytest.raises(CorruptFileError, match="corrupt or invalid"):
            parse_document(corrupt_file)

    def test_invalid_utf8_txt(self, tmp_path):
        """Text file with non-UTF-8 bytes should raise EncodingError."""
        invalid_txt = tmp_path / "invalid_encoding.txt"
        # Write invalid UTF-8 sequence
        invalid_txt.write_bytes(b"Synthetic header \x80\x81\xFF invalid bytes")
        with pytest.raises(EncodingError, match="Failed to decode text file.*as UTF-8"):
            parse_document(invalid_txt)


# ==============================================================================
# Test Cases: File Existence, Format Rejection, and Size Limits
# ==============================================================================


class TestValidationAndLimits:
    """Tests verifying file existence, unsupported extensions, and file size limits."""

    def test_missing_file(self, tmp_path):
        """Non-existent file should raise FileNotFoundParserError."""
        missing = tmp_path / "nonexistent_resume.pdf"
        with pytest.raises(FileNotFoundParserError, match="not found"):
            parse_document(missing)

    @pytest.mark.parametrize(
        "unsupported_ext",
        [".jpg", ".png", ".exe", ".rtf", ".odt", ".csv", ".json"],
    )
    def test_unsupported_extensions(self, tmp_path, unsupported_ext):
        """Unsupported file extensions should raise UnsupportedFormatError."""
        bad_file = tmp_path / f"test{unsupported_ext}"
        bad_file.write_text("synthetic content", encoding="utf-8")
        with pytest.raises(UnsupportedFormatError, match="Unsupported file extension"):
            parse_document(bad_file)

    def test_file_size_limit_exceeded(self, tmp_path):
        """File exceeding configured max_file_size_bytes should raise FileTooLargeError."""
        txt_file = create_synthetic_txt(tmp_path / "large.txt", "A" * 2048)
        # Configure parser with 1024 bytes limit
        parser = DocumentParser(max_file_size_bytes=1024)
        with pytest.raises(FileTooLargeError, match="exceeds maximum allowed limit"):
            parser.parse(txt_file)

    def test_file_size_limit_permitted(self, tmp_path):
        """File within configured size limit should succeed."""
        txt_file = create_synthetic_txt(tmp_path / "normal.txt", "A" * 500)
        parser = DocumentParser(max_file_size_bytes=1024)
        result = parser.parse(txt_file)
        assert result.character_count == 500
